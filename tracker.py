import csv
import json
import os
import time
from datetime import datetime, timezone

import requests
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

PLAYERS_FILE = "players.json"
HISTORY_FILE = "prices.csv"
SIGNALS_FILE = "signals.csv"
REPORT_FILE = "report.md"
XLSX_FILE = "market_tracker.xlsx"

API_BASE = "https://api.parse.bot/scraper/a1271aad-bcbf-4464-8762-47f1d15efa81"

WATCHLIST = [
    "Kylian Mbappe", "Erling Haaland", "Vinicius Junior", "Lamine Yamal",
    "Jude Bellingham", "Bukayo Saka", "Mohamed Salah", "Harry Kane",
    "Jamal Musiala", "Theo Hernandez", "Nico Schlotterbeck", "Cole Palmer",
    "Pedri", "Rodri", "Federico Valverde", "Achraf Hakimi", "Alphonso Davies",
    "Florian Wirtz", "Alexander Isak", "Khvicha Kvaratskhelia",
]

HEADERS = {
    "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 Chrome/153 Safari/537.36",
    "Accept": "application/json",
}


def api_get(endpoint, params):
    r = requests.get(f"{API_BASE}/{endpoint}", params=params, headers=HEADERS, timeout=60)
    r.raise_for_status()
    data = r.json()
    if isinstance(data, dict) and data.get("status") not in (None, "success"):
        raise RuntimeError(str(data))
    return data.get("data", data) if isinstance(data, dict) else data


def normalize(s):
    return "".join(c.lower() for c in str(s).strip() if c.isalnum())


def find_player_id(name):
    data = api_get("search_players_fc27", {"name": name, "page": 1})
    results = data.get("results", []) if isinstance(data, dict) else []
    if not results:
        # Some API revisions use query instead of name.
        data = api_get("search_players_fc27", {"query": name, "page": 1})
        results = data.get("results", []) if isinstance(data, dict) else []

    if not results:
        return None

    target = normalize(name)
    exact = [p for p in results if normalize(p.get("name", "")) == target]
    candidate = exact[0] if exact else results[0]
    return int(candidate["id"])


def load_players():
    if os.path.exists(PLAYERS_FILE):
        with open(PLAYERS_FILE, encoding="utf-8") as f:
            players = json.load(f)
    else:
        players = [{"name": n} for n in WATCHLIST]

    by_name = {p["name"]: p for p in players}
    return [by_name.get(n, {"name": n}) for n in WATCHLIST]


def snapshot(ids):
    params = {
        "player_ids": ",".join(str(x) for x in ids),
        "platform": "ps",
    }
    return api_get("get_fc27_market_snapshot", params)


def append_history(rows):
    exists = os.path.exists(HISTORY_FILE) and os.path.getsize(HISTORY_FILE) > 0
    fields = ["timestamp_utc", "player", "futbin_id", "platform", "price",
              "previous_price", "change_coins", "change_percent"]
    with open(HISTORY_FILE, "a", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=fields)
        if not exists:
            w.writeheader()
        w.writerows(rows)


def load_history():
    if not os.path.exists(HISTORY_FILE):
        return []
    with open(HISTORY_FILE, newline="", encoding="utf-8") as f:
        return list(csv.DictReader(f))


def pct(a, b):
    if a in (None, "") or b in (None, ""):
        return None
    try:
        a, b = float(a), float(b)
        return round((a - b) / b * 100, 2) if b else None
    except (TypeError, ValueError):
        return None


def build_signals(rows):
    by_player = {}
    for r in rows:
        by_player.setdefault(r["player"], []).append(r)

    out = []
    for name, rs in by_player.items():
        rs.sort(key=lambda x: x["timestamp_utc"])
        current = int(rs[-1]["price"])
        previous = int(rs[-2]["price"]) if len(rs) >= 2 else None
        one_day = None
        seven_day = None
        now = datetime.fromisoformat(rs[-1]["timestamp_utc"].replace("Z", "+00:00"))
        for r in reversed(rs[:-1]):
            t = datetime.fromisoformat(r["timestamp_utc"].replace("Z", "+00:00"))
            age_h = (now - t).total_seconds() / 3600
            if one_day is None and age_h >= 24:
                one_day = int(r["price"])
            if seven_day is None and age_h >= 168:
                seven_day = int(r["price"])
                break

        ch = pct(current, previous)
        ch24 = pct(current, one_day)
        ch7 = pct(current, seven_day)
        if ch is not None and ch >= 3:
            signal = "RISING"
        elif ch is not None and ch <= -3:
            signal = "FALLING"
        else:
            signal = "STABLE"

        out.append({
            "player": name,
            "platform": "PS/Xbox",
            "price": current,
            "change_2h_pct": ch,
            "change_24h_pct": ch24,
            "change_7d_pct": ch7,
            "signal": signal,
            "snapshots": len(rs),
        })
    return sorted(out, key=lambda x: (x["change_24h_pct"] is None, -(x["change_24h_pct"] or 0)))


def write_signals(signals):
    fields = ["player", "platform", "price", "change_2h_pct", "change_24h_pct",
              "change_7d_pct", "signal", "snapshots"]
    with open(SIGNALS_FILE, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=fields)
        w.writeheader()
        w.writerows(signals)


def write_report(signals):
    lines = [
        "# FC27 Market Report",
        "",
        f"Updated: {datetime.now(timezone.utc).isoformat()}",
        "",
        "Prices are PlayStation/console market snapshots. Changes are calculated from this tracker's stored history.",
        "",
        "| Player | Price | 2h | 24h | 7d | Signal |",
        "|---|---:|---:|---:|---:|---|",
    ]
    for s in signals:
        fmt = lambda x: "—" if x is None else f"{x:+.2f}%"
        lines.append(f"| {s['player']} | {s['price']:,} | {fmt(s['change_2h_pct'])} | {fmt(s['change_24h_pct'])} | {fmt(s['change_7d_pct'])} | {s['signal']} |")
    with open(REPORT_FILE, "w", encoding="utf-8") as f:
        f.write("\n".join(lines) + "\n")


def write_xlsx(players, history, signals):
    if os.path.exists(XLSX_FILE):
        wb = load_workbook(XLSX_FILE)
    else:
        wb = Workbook()
        wb.remove(wb.active)

    def sheet(name):
        return wb[name] if name in wb.sheetnames else wb.create_sheet(name)

    ws = sheet("Dashboard")
    ws.delete_rows(1, ws.max_row)
    ws.append(["FC27 MARKET TRACKER"])
    ws.append(["Last update UTC", datetime.now(timezone.utc).isoformat()])
    ws.append([])
    ws.append(["Player", "Price", "2h %", "24h %", "7d %", "Signal", "Snapshots"])
    for s in signals:
        ws.append([s["player"], s["price"], s["change_2h_pct"], s["change_24h_pct"], s["change_7d_pct"], s["signal"], s["snapshots"]])

    hp = sheet("Price History")
    hp.delete_rows(1, hp.max_row)
    if history:
        fields = list(history[0].keys())
        hp.append(fields)
        for r in history:
            hp.append([r.get(k, "") for k in fields])

    sp = sheet("Signals")
    sp.delete_rows(1, sp.max_row)
    if signals:
        fields = list(signals[0].keys())
        sp.append(fields)
        for r in signals:
            sp.append([r.get(k, "") for k in fields])

    pp = sheet("Players")
    pp.delete_rows(1, pp.max_row)
    pp.append(["Player", "FUTBIN ID"])
    for p in players:
        pp.append([p["name"], p.get("futbin_id", "")])

    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for col in range(1, ws.max_column + 1):
            ws.column_dimensions[get_column_letter(col)].width = min(28, max(12, max(len(str(ws.cell(r, col).value or "")) for r in range(1, min(ws.max_row, 50) + 1)) + 2))
        for row in ws.iter_rows():
            for cell in row:
                cell.alignment = Alignment(vertical="center")
    wb.save(XLSX_FILE)


def main():
    players = load_players()

    # Resolve missing card IDs once, then persist them.
    for p in players:
        if not p.get("futbin_id"):
            try:
                p["futbin_id"] = find_player_id(p["name"])
                print(f"Resolved {p['name']}: {p['futbin_id']}")
                time.sleep(0.4)
            except Exception as e:
                print(f"ERROR resolving {p['name']}: {e}")

    valid = [p for p in players if p.get("futbin_id")]
    if not valid:
        raise SystemExit("No player IDs could be resolved from the external API")

    ids = [int(p["futbin_id"]) for p in valid]
    data = snapshot(ids)
    price_rows = data.get("players", []) if isinstance(data, dict) else []

    price_map = {int(x["player_id"]): x.get("price") for x in price_rows if x.get("player_id") is not None}
    timestamp_ms = data.get("timestamp") if isinstance(data, dict) else None
    timestamp = datetime.fromtimestamp(timestamp_ms / 1000, tz=timezone.utc).isoformat() if timestamp_ms else datetime.now(timezone.utc).isoformat()

    history = load_history()
    previous = {}
    for r in history:
        previous[(r["player"], r["platform"])] = r["price"]

    rows = []
    for p in valid:
        pid = int(p["futbin_id"])
        price = price_map.get(pid)
        if price is None or int(price) <= 0:
            print(f"ERROR {p['name']}: no current price")
            continue
        price = int(price)
        prev = previous.get((p["name"], "PS/Xbox"))
        change = price - int(prev) if prev is not None else None
        change_pct = pct(price, prev)
        print(f"OK {p['name']}: {price:,} coins")
        rows.append({
            "timestamp_utc": timestamp,
            "player": p["name"],
            "futbin_id": pid,
            "platform": "PS/Xbox",
            "price": price,
            "previous_price": prev,
            "change_coins": change,
            "change_percent": change_pct,
        })

    if not rows:
        raise SystemExit("No prices were collected from the external API")

    with open(PLAYERS_FILE, "w", encoding="utf-8") as f:
        json.dump(players, f, indent=2, ensure_ascii=False)
        f.write("\n")

    append_history(rows)
    history = load_history()
    signals = build_signals(history)
    write_signals(signals)
    write_report(signals)
    write_xlsx(players, history, signals)
    print(f"Collected prices for {len(rows)} players and updated {XLSX_FILE}")


if __name__ == "__main__":
    main()
